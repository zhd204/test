from __future__ import annotations

import json
import shutil
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


WORKSPACE = Path(__file__).resolve().parents[1]
SOURCE = WORKSPACE / "findings .xlsx"
OUTPUT_DIR = WORKSPACE / "outputs" / "unique-findings"
OUTPUT = OUTPUT_DIR / "findings_with_unique_findings.xlsx"
SHEET_NAME = "Unique Findings"

# Source column positions (0-based) within columns A..G of the findings sheet.
FINDING_NAME_COL = 0  # A - Finding Name
SEVERITY_COL = 5  # F - Severity
SOLUTION_COL = 6  # G - Solution

# Highest risk first; any unrecognised value sorts after these.
SEVERITY_ORDER = ["Critical", "High", "Medium", "Low", "Info", "Informational", "None"]


def collect_unique_findings():
    if not SOURCE.exists():
        raise FileNotFoundError(
            f"Source workbook not found: {SOURCE}\n"
            "Place 'findings .xlsx' in the folder one level above this script."
        )
    workbook = load_workbook(SOURCE, read_only=True, data_only=False)
    source_sheet = workbook["findings "]

    counts = Counter()
    solutions = defaultdict(Counter)
    severities = defaultdict(Counter)
    for row in source_sheet.iter_rows(
        min_row=2,
        min_col=1,
        max_col=7,
        values_only=True,
    ):
        finding_name = row[FINDING_NAME_COL]
        severity = row[SEVERITY_COL]
        solution = row[SOLUTION_COL]
        if finding_name is None or str(finding_name).strip() == "":
            continue
        name = str(finding_name).strip()
        counts[name] += 1
        if solution is not None and str(solution).strip():
            solutions[name][str(solution).strip()] += 1
        if severity is not None and str(severity).strip():
            severities[name][str(severity).strip()] += 1

    workbook.close()
    return counts, solutions, severities


def severity_rank(severity):
    lowered = [level.lower() for level in SEVERITY_ORDER]
    try:
        return lowered.index(severity.lower())
    except ValueError:
        return len(lowered)


def combine_severities(severity_counts):
    """Return the finding's severity; list all distinct values, highest first, if they differ."""
    if not severity_counts:
        return "Not specified"
    ranked = sorted(severity_counts, key=lambda value: (severity_rank(value), value))
    return ", ".join(ranked)


def combine_solutions(solution_counts):
    ranked = sorted(solution_counts.items(), key=lambda item: (-item[1], item[0]))
    if not ranked:
        return "No solution provided in source data."
    if len(ranked) == 1:
        return ranked[0][0]

    combined = "\n\n".join(
        f"Solution variant {index}:\n{solution}"
        for index, (solution, _) in enumerate(ranked, start=1)
    )
    if len(combined) > 32767:
        raise ValueError(
            "A combined solution exceeds Excel's 32,767-character cell limit."
        )
    return combined


def build_workbook():
    counts, solutions, severities = collect_unique_findings()
    rows = sorted(counts.items(), key=lambda item: (-item[1], item[0]))

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)

    workbook = load_workbook(OUTPUT)
    if SHEET_NAME in workbook.sheetnames:
        del workbook[SHEET_NAME]
    sheet = workbook.create_sheet(SHEET_NAME)
    sheet.sheet_properties.tabColor = "1F4E78"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"

    headers = ["Finding Name", "Severity", "Solution", "Count"]
    sheet.append(headers)
    for finding_name, count in rows:
        sheet.append(
            [
                finding_name,
                combine_severities(severities[finding_name]),
                combine_solutions(solutions[finding_name]),
                count,
            ]
        )

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10, color="222222")

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 24

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=4):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (1, 3))
        row[1].alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        row[3].number_format = "#,##0"
        row[3].alignment = Alignment(horizontal="right", vertical="top")
        sheet.row_dimensions[row[0].row].height = 60

    sheet.column_dimensions["A"].width = 72
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 100
    sheet.column_dimensions["D"].width = 12

    table_ref = f"A1:D{sheet.max_row}"
    table = Table(displayName="UniqueFindingsTable", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)

    workbook.save(OUTPUT)
    workbook.close()

    validate_output(counts)


def validate_output(counts):
    workbook = load_workbook(OUTPUT, read_only=True, data_only=False)
    source_sheet = workbook["findings "]
    unique_sheet = workbook[SHEET_NAME]

    output_counts = {}
    for finding_name, _, _, count in unique_sheet.iter_rows(min_row=2, values_only=True):
        output_counts[finding_name] = count

    expected_total = sum(counts.values())
    actual_total = sum(output_counts.values())
    if output_counts != dict(counts):
        raise AssertionError("Saved unique-finding counts do not match the source aggregation.")
    if actual_total != expected_total:
        raise AssertionError("Saved counts do not reconcile to the source row total.")
    if source_sheet.max_row != 110807 or source_sheet.max_column != 12:
        raise AssertionError("The original findings sheet dimensions changed unexpectedly.")

    result = {
        "output": str(OUTPUT),
        "sheets": workbook.sheetnames,
        "source_dimensions": [source_sheet.max_row, source_sheet.max_column],
        "unique_findings_rows": unique_sheet.max_row - 1,
        "source_rows_counted": actual_total,
        "top_five": list(unique_sheet.iter_rows(min_row=2, max_row=6, values_only=True)),
    }
    workbook.close()
    print(json.dumps(result, indent=2, default=str))


if __name__ == "__main__":
    build_workbook()
